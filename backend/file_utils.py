import os
import uuid
import zipfile
from pathlib import Path
from typing import List
from openpyxl import load_workbook, Workbook

UPLOAD_DIR = os.environ.get("UPLOAD_DIR", "/data/uploads")
OUTPUT_DIR = os.environ.get("OUTPUT_DIR", "/data/outputs")

EXAM_EXTENSIONS = {
    "word": [".docx"],
    "excel": [".xlsx", ".xls"],
    "powerpoint": [".pptx"]
}

CRITERIA_EXTENSION = ".xlsx"

def ensure_dirs():
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    os.makedirs(OUTPUT_DIR, exist_ok=True)

def detect_category(filename: str):
    ext = Path(filename).suffix.lower()
    for cat, exts in EXAM_EXTENSIONS.items():
        if ext in exts:
            return cat
    return None

def is_criteria_file(filename: str) -> bool:
    return Path(filename).suffix.lower() == CRITERIA_EXTENSION

def save_upload(file_content: bytes, filename: str, subfolder: str) -> str:
    folder = os.path.join(UPLOAD_DIR, subfolder)
    os.makedirs(folder, exist_ok=True)
    unique_name = f"{uuid.uuid4().hex}_{filename}"
    path = os.path.join(folder, unique_name)
    with open(path, "wb") as f:
        f.write(file_content)
    return path

def delete_file(path: str):
    if os.path.exists(path):
        os.remove(path)

def merge_criteria_files(criteria_paths: List[str], output_path: str) -> str:
    merged_wb = Workbook()
    merged_ws = merged_wb.active
    merged_ws.title = "Merged Grading Criteria"

    header_written = False
    row_offset = 0

    for idx, cpath in enumerate(criteria_paths):
        if not os.path.exists(cpath):
            continue
        wb = load_workbook(cpath, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()

        if not rows:
            continue

        source_name = Path(cpath).stem

        if not header_written:
            header_row = rows[0]
            data_rows = rows[1:]
        else:
            header_row = None
            data_rows = rows

        if header_row is not None:
            merged_ws.cell(row=1, column=1, value="Source")
            for col_idx, val in enumerate(header_row, 2):
                merged_ws.cell(row=1, column=col_idx, value=val)
            header_written = True

        for row in data_rows:
            row_offset += 1
            merged_ws.cell(row=row_offset + 1, column=1, value=f"[{source_name}]")
            for col_idx, val in enumerate(row):
                merged_ws.cell(row=row_offset + 1, column=col_idx + 2, value=val)

    if row_offset == 0 and not header_written:
        merged_ws.cell(row=1, column=1, value="No grading criteria found")

    merged_wb.save(output_path)
    merged_wb.close()
    return output_path

def package_exam_set(
    exam_paths: List[str],
    exam_names: List[str],
    criteria_path: str,
    output_dir: str
) -> str:
    os.makedirs(output_dir, exist_ok=True)
    zip_name = f"exam_set_{uuid.uuid4().hex[:8]}.zip"
    zip_path = os.path.join(output_dir, zip_name)

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for ep, en in zip(exam_paths, exam_names):
            if os.path.exists(ep):
                zf.write(ep, en)
        if os.path.exists(criteria_path):
            zf.write(criteria_path, "Final_Grading_Sheet.xlsx")

    return zip_path
